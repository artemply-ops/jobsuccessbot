"""Извлечение текста из загруженных файлов."""

import io
from pathlib import Path

MAX_CHARS = 60_000
SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".rtf"}


class ExtractError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        if not text.strip():
            raise ExtractError("В PDF нет текстового слоя (скан). Пришлите DOCX или текст.")
    elif ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif ext in {".txt", ".md", ".rtf"}:
        text = _decode(data)
    else:
        raise ExtractError(f"Формат {ext or 'без расширения'} не поддерживается. Можно: PDF, DOCX, TXT.")
    text = "\n".join(line.rstrip() for line in text.splitlines() if line.strip())
    return text[:MAX_CHARS]


def _decode(data: bytes) -> str:
    # Windows-файлы часто в CP1251
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")
