"""Выгрузка пакета в DOCX и PDF. Здесь и только здесь обезличенные маркеры меняются на реальные данные."""

import io
from pathlib import Path

from docx import Document
from docx.shared import Pt
from fpdf import FPDF

from app.models import Contact, Package
from app.services.pii import Vault, restore

FONT_CANDIDATES = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
]


class FlagsNotAccepted(ValueError):
    pass


def _flagged(pkg: Package) -> set[str]:
    return {f["text"].strip() for f in pkg.fact_flags or []}


def cv_view(pkg: Package, contact: Contact | None, drop_flagged: bool = True) -> dict:
    """CV с реальными контактами. Неподтверждённые пункты CV выкидываем."""
    flagged = _flagged(pkg) if drop_flagged else set()
    vault = Vault(dict(contact.vault or {})) if contact else Vault()
    r = lambda t: restore(t, contact.full_name if contact else None, contact.email if contact else None,  # noqa: E731
                          contact.phone if contact else None, vault)
    cv = pkg.cv
    return {
        "name": (contact.full_name if contact else None) or "",
        "contacts": " · ".join(x for x in [contact.email, contact.phone, contact.city] if x) if contact else "",
        "headline": r(cv["headline"]),
        "summary": r(cv["summary"]),
        "skills": [r(s) for s in cv["skills"]],
        "experience": [{**e, "bullets": [r(b) for b in e["bullets"] if b.strip() not in flagged]}
                       for e in cv["experience"]],
        "education": [r(x) for x in cv.get("education", [])],
        "languages": [r(x) for x in cv.get("languages", [])],
        "cover_letter": r(pkg.cover_letter),
        "recruiter_message": r(pkg.recruiter_message),
    }


def _check(pkg: Package, accept_flags: bool) -> None:
    letter_flags = [f for f in pkg.fact_flags or [] if f.get("where") != "cv"]
    if letter_flags and not accept_flags:
        raise FlagsNotAccepted("В письме есть утверждения, не подтверждённые профилем. "
                               "Исправьте или подтвердите их явно.")


def to_docx(pkg: Package, contact: Contact | None, accept_flags: bool = False) -> bytes:
    _check(pkg, accept_flags)
    v = cv_view(pkg, contact)
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(v["name"] or v["headline"], level=0)
    if v["name"]:
        doc.add_paragraph(v["headline"]).runs[0].bold = True
    if v["contacts"]:
        doc.add_paragraph(v["contacts"])
    doc.add_paragraph(v["summary"])
    doc.add_heading("Ключевые навыки", level=2)
    doc.add_paragraph(", ".join(v["skills"]))
    doc.add_heading("Опыт работы", level=2)
    for e in v["experience"]:
        p = doc.add_paragraph()
        p.add_run(f"{e['role']} — {e['company']}").bold = True
        p.add_run(f"   {e['period']}")
        for b in e["bullets"]:
            doc.add_paragraph(b, style="List Bullet")
    if v["education"]:
        doc.add_heading("Образование", level=2)
        for x in v["education"]:
            doc.add_paragraph(x)
    if v["languages"]:
        doc.add_heading("Языки", level=2)
        doc.add_paragraph(", ".join(v["languages"]))
    doc.add_page_break()
    doc.add_heading("Сопроводительное письмо", level=1)
    for para in v["cover_letter"].split("\n"):
        if para.strip():
            doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _font() -> tuple[str, str]:
    for reg, bold in FONT_CANDIDATES:
        if Path(reg).exists():
            return reg, bold if Path(bold).exists() else reg
    raise RuntimeError("Не найден TTF-шрифт с кириллицей (Arial или DejaVu). Установите fonts-dejavu-core.")


def to_pdf(pkg: Package, contact: Contact | None, accept_flags: bool = False) -> bytes:
    _check(pkg, accept_flags)
    v = cv_view(pkg, contact)
    reg, bold = _font()
    pdf = FPDF(format="A4")
    pdf.set_margins(18, 16, 18)
    pdf.add_font("Main", "", reg)
    pdf.add_font("Main", "B", bold)
    pdf.add_page()
    w = pdf.epw

    def h(text: str, size: int = 13):
        pdf.ln(2)
        pdf.set_font("Main", "B", size)
        pdf.multi_cell(w, 7, text)
        pdf.set_font("Main", "", 10.5)

    def p(text: str, line: float = 5.2):
        pdf.set_font("Main", "", 10.5)
        pdf.multi_cell(w, line, text)

    h(v["name"] or v["headline"], 18)
    if v["name"]:
        h(v["headline"], 12)
    if v["contacts"]:
        p(v["contacts"])
    pdf.ln(1)
    p(v["summary"])
    h("Ключевые навыки")
    p(", ".join(v["skills"]))
    h("Опыт работы")
    for e in v["experience"]:
        pdf.set_font("Main", "B", 11)
        pdf.multi_cell(w, 6, f"{e['role']} — {e['company']}   ({e['period']})")
        for b in e["bullets"]:
            p(f"•  {b}")
        pdf.ln(1)
    if v["education"]:
        h("Образование")
        for x in v["education"]:
            p(x)
    if v["languages"]:
        h("Языки")
        p(", ".join(v["languages"]))
    pdf.add_page()
    h("Сопроводительное письмо", 14)
    for para in v["cover_letter"].split("\n"):
        if para.strip():
            p(para)
            pdf.ln(1.5)
    return bytes(pdf.output())
